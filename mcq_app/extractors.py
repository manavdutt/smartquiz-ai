import fitz  # PyMuPDF
import re

# Word files (.docx) read karne ke liye module import karein
# Note: Iske liye aapke environment me 'python-docx' installed hona chahiye (pip install python-docx)
import docx


# ─────────────────────────────────────────────────────────────
#  1. WORD FILE (.DOCX) TEXT EXTRACTOR
# ─────────────────────────────────────────────────────────────
def extract_from_docx(docx_file):
    """Django uploaded Word (.docx) file se cleanly text nikalne ke liye"""
    try:
        docx_file.seek(0)  # Safe pointer reset
        doc = docx.Document(docx_file)
        full_text = []

        # Har paragraph ka text uthao
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)

        # Agar file me Tables hain, toh unka data bhi nikal lo
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))

        docx_file.seek(0)  # Reset pointer back for Django
        return "\n".join(full_text).strip()
    except Exception as e:
        raise ValueError(f"Could not read Word file: {str(e)}")


# ─────────────────────────────────────────────────────────────
#  2. PDF TEXT EXTRACTOR
# ─────────────────────────────────────────────────────────────
def extract_from_pdf(pdf_file):
    text = ""
    try:
        pdf_file.seek(0)
        pdf_bytes = pdf_file.read()
        pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")

        try:
            for page in pdf_document:
                text += page.get_text()
        finally:
            pdf_document.close()  # ← Error ho ya na ho — hamesha close hoga ✅

        pdf_file.seek(0)
    except Exception as e:
        raise ValueError(f"Could not read PDF file: {str(e)}")

    return text.strip()


# ─────────────────────────────────────────────────────────────
#  3. YOUTUBE LINK EXTRACTOR
# ─────────────────────────────────────────────────────────────
def get_youtube_id(url):
    """Har tarah ke YouTube URLs se 11 digit ki Video ID nikalne ke liye"""
    patterns = [
        r'v=([a-zA-Z0-9_-]{11})',
        r'youtu\.be/([a-zA-Z0-9_-]{11})',
        r'shorts/([a-zA-Z0-9_-]{11})',
        r'embed/([a-zA-Z0-9_-]{11})',
    ]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return None


def extract_from_youtube(url):
    """YouTube video ki transcript automatic best language mein nikalne ke liye"""
    from youtube_transcript_api import YouTubeTranscriptApi

    video_id = get_youtube_id(url)
    if not video_id:
        raise ValueError("Invalid YouTube URL! Please enter a valid link.")

    target_languages = ['en', 'hi', 'ar', 'es', 'fr', 'zh', 'pt']

    api = YouTubeTranscriptApi()

    try:
        fetched = api.fetch(video_id, languages=target_languages)
        text = " ".join([entry.text for entry in fetched])
        return text.strip()
    except Exception:
        try:
            fetched = api.fetch(video_id)
            text = " ".join([entry.text for entry in fetched])
            return text.strip()
        except Exception as e:
            raise ValueError(f"Transcript not available for this video: {str(e)}")
